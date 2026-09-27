"""Generate Panduan Pengguna (User Manual) FleetVision AI (.docx). Usage: python docs/generate_user_manual_docx.py"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def h1(t):
    doc.add_heading(t, level=1)


def h2(t):
    doc.add_heading(t, level=2)


def p(t, bold=False):
    par = doc.add_paragraph()
    run = par.add_run(t)
    run.bold = bold
    return par


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def numbered(items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def tip(t):
    par = doc.add_paragraph()
    run = par.add_run("Tips: " + t)
    run.italic = True
    return par


def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        t.cell(0, i).text = htxt
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    doc.add_paragraph()
    return t


# ================= COVER =================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Panduan Pengguna\nFleetVision AI")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Manual pemakaian aplikasi untuk Admin, Sopir, dan Publik\nhttps://dtrack.yusufandrika.dev")
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(0x47, 0x5B, 0x69)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Real-Time Fleet Tracking & Driver Monitoring System")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

doc.add_page_break()

# ================= 1 =================
h1("1. Pengenalan Aplikasi")
p("FleetVision AI adalah aplikasi pemantauan armada logistik secara real-time. "
  "Posisi kendaraan terpantau live di peta, kondisi sopir dipantau kamera AI "
  "(sabuk pengaman, kantuk, main HP, fokus pandangan), dan setiap perjalanan tercatat history-nya.")
h2("1.1 Tiga peran pengguna")
table(["Peran", "Kegunaan", "Alamat"],
      [["Admin", "Kelola armada, buat rute trip, pantau semua unit + kamera live, lihat alert",
        "/login lalu otomatis ke /admin"],
       ["Sopir", "Menerima tugas rute, menjalankan trip dengan GPS + kamera AI, melihat riwayat sendiri",
        "/login lalu otomatis ke /driver"],
       ["Publik", "Melacak posisi kendaraan cukup dengan nomor plat, tanpa login", "/track"]])
h2("1.2 Masuk (login)")
numbered(["Buka alamat aplikasi, pilih Masuk bila masih di halaman pelacakan.",
          "Isi email dan password, tekan Masuk Dashboard.",
          "Admin otomatis masuk ke dashboard admin, sopir ke dashboard sopir."])
tip("Akun diberikan oleh admin. Bila lupa password, hubungi admin.")

# ================= 2 ADMIN =================
h1("2. Panduan Admin")
h2("2.1 Dashboard")
p("Halaman utama berisi kartu ringkasan: jumlah kendaraan, sopir, trip hari ini, dan alert. "
  "Gunakan sebagai pantauan sekilas sebelum masuk ke menu detail.")

h2("2.2 Armada (Fleet)")
bullets(["Menu Armada menampilkan peta semua kendaraan beserta daftarnya.",
         "Klik kendaraan di daftar/peta untuk membuka detail unit: posisi terakhir, kecepatan, status.",
         "Di detail unit ada panel Live Camera (tayangan kamera sopir, bila sedang trip) dan "
         "AI Behavior Detection: Face Detected, Seatbelt On, Eyes Open, No Fatigue, No Phone, Looking Ahead "
         "plus angka Eye Closure, Yaw, dan Pitch."])
tip("Kamera live hanya tampil saat sopir sedang menjalankan trip dan kameranya aktif.")

h2("2.3 Membuat trip + rute")
p("Admin menentukan titik awal dan tujuan perjalanan untuk sopir:")
numbered(["Buka menu Trip, tekan Buat Trip.",
          "Pilih Kendaraan dan Driver.",
          "Isi Titik awal dan Tujuan (nama tempat, contoh: Jakarta, Surabaya).",
          "Tentukan koordinat dengan salah satu cara: ketik nama kota/alamat di Cari lokasi lalu klik hasil, "
          "atau pilih titik Awal/Tujuan lalu klik langsung di peta.",
          "Estimasi jarak terisi otomatis. Tekan Simpan Trip Terjadwal."])
tip("Untuk nama perusahaan yang tidak ketemu saat dicari (contoh nama PT), tambah nama kota di pencarian "
    "atau klik langsung lokasinya di peta.")

h2("2.4 Melihat history perjalanan")
p("Setiap trip — baik yang dibuat admin maupun yang dimulai sendiri oleh sopir — tercatat history-nya:")
bullets(["Halaman Trip langsung menampilkan panel History Perjalanan untuk trip terbaru: "
         "peta rute rencana (garis biru putus-putus) + jejak jalan aktual (garis hijau) + marker awal, tujuan, "
         "dan posisi terakhir.",
         "Kartu angka: Estimasi rencana, Sudah jalan (km), Sisa ke tujuan, dan Durasi.",
         "Untuk melihat trip lain, tekan tombol History pada baris trip di tabel."])
tip("Trip yang masih Terjadwal belum punya jejak GPS. Garis lurus antar-pulau (mis. Jawa-Kalimantan) "
    "adalah jalur penyeberangan laut, bukan jalan darat.")

h2("2.5 Driver, Alert, dan Settings")
bullets(["Menu Driver: daftar sopir, klik untuk detail, tugaskan kendaraan bila diperlukan.",
         "Menu Trip: riwayat semua perjalanan dengan filter Semua/Aktif/Terjadwal/Selesai/Batal.",
         "Menu Alerts: daftar peringatan (sabuk, kantuk, HP, ngebut, durasi mengemudi). Tandai dibaca bila sudah ditangani.",
         "Menu Settings: pengaturan aplikasi."])

# ================= 3 SOPIR =================
h1("3. Panduan Sopir")
h2("3.1 Dashboard sopir")
p("Halaman utama sopir berisi sapaan, kartu Unit/Status/Durasi/Speed/Jarak/AI Cam, peta posisi, dan profil. "
  "Kartu Jarak menampilkan jarak rencana dari admin; jarak yang sudah ditempuh tampil di kartu perjalanan.")

h2("3.2 Mengerjakan tugas rute dari admin")
numbered(["Bila ada kartu Tugas rute dari admin, periksa asal, tujuan, dan estimasinya.",
          "Di lokasi awal, tekan Mulai Rute Ini dan izinkan Lokasi serta Kamera bila diminta.",
          "Banner rute menampilkan asal-tujuan, Estimasi, Sisa, dan Ditempuh.",
          "Peta menampilkan marker tujuan dan garis rute. Geser peta untuk melihat rute (mode Bebas); "
          "tekan tombol Mengikuti untuk menempel ke posisi kendaraan lagi.",
          "Setelah sampai tujuan, tekan Selesaikan Trip."])

h2("3.3 Jalan bebas (tanpa tugas admin)")
p("Bila tidak ada tugas rute, tekan Mulai Trip untuk jalan bebas. GPS dan kamera tetap aktif dan "
  "perjalanan tetap tercatat di history.")

h2("3.4 Indikator kamera AI")
p("Selama trip, kamera memantau dan menampilkan: Wajah, Seatbelt, Mata Terbuka, Tidak Lelah, Tanpa HP, "
  "Fokus Depan, plus persentase mata dan sudut kepala (Yaw/Pitch). Pastikan wajah terlihat jelas di preview "
  "kamera sendiri dan pakai sabuk pengaman agar tidak muncul peringatan.")

h2("3.5 Riwayat dan alert saya")
bullets(["Menu Trip Saya: daftar perjalanan sendiri lengkap dengan rute, status, dan jarak. "
         "Klik salah satu untuk membuka detail peta history-nya.",
         "Menu Alerts: peringatan yang ditujukan ke Anda."])

# ================= 4 PUBLIK =================
h1("4. Panduan Publik (Tanpa Login)")
numbered(["Buka halaman Lacak armada.",
          "Masukkan nomor plat kendaraan (contoh: B-1234-ABC), tekan Lacak.",
          "Terlihat posisi live, kecepatan, dan waktu update. Bila kendaraan sedang bertrip, terlihat juga "
          "tujuan, estimasi jarak, sisa jarak, dan garis rutenya.",
          "Tekan Masuk di pojok atas bila ingin login sebagai admin/sopir."])

# ================= 5 ISTILAH =================
h1("5. Arti Status dan Tanda di Peta")
h2("5.1 Status trip")
table(["Status", "Artinya"],
      [["Terjadwal (PLANNED)", "Rute dibuat admin, sopir belum jalan"],
       ["Aktif (IN_PROGRESS)", "Sopir sedang dalam perjalanan"],
       ["Selesai (COMPLETED)", "Perjalanan sudah ditutup dengan jarak akhir"],
       ["Batal (CANCELLED)", "Trip dibatalkan"]])
h2("5.2 Tanda di peta")
table(["Tanda", "Artinya"],
      [["Marker truk hijau/biru", "Posisi kendaraan"],
       ["Marker pin hijau", "Titik awal"],
       ["Marker target merah", "Tujuan"],
       ["Garis biru putus-putus", "Rute rencana"],
       ["Garis hijau", "Jejak jalan yang sudah dilalui"]])

# ================= 6 FAQ =================
h1("6. Pertanyaan Umum")
table(["Pertanyaan", "Jawaban"],
      [["Kenapa diminta izin lokasi/kamera?", "Wajib untuk GPS tracking dan monitoring AI. Tanpa izin, trip tidak bisa dimulai."],
       ["Kamera saya hitam di preview?", "Periksa izin kamera di browser/HP, pastikan tidak dipakai aplikasi lain, lalu mulai ulang trip."],
       ["Jarak masih 0?", "Kartu Jarak memakai jarak rencana admin. Bila trip tanpa rencana, Jarak mengikuti jarak tempuh yang bertambah sambil jalan."],
       ["Peta bergerak sendiri?", "Itu mode Mengikuti. Geser peta untuk mode Bebas."],
       ["History kosong?", "Trip terjadwal memang belum ada jejak. Untuk trip yang sudah jalan, jejak muncul bila GPS terkirim ke server."],
       ["Tidak bisa login?", "Periksa email/password ke admin. Bila halaman error, refresh dan coba lagi."]])

doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = end.add_run("— Selesai —  Panduan Pengguna FleetVision AI")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

doc.save("Panduan-Pengguna-FleetVision-AI.docx")
print("saved Panduan-Pengguna-FleetVision-AI.docx")
