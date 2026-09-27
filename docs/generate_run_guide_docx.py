"""Generate Panduan Cara Menjalankan Aplikasi FleetVision AI (.docx). Usage: python docs/generate_run_guide_docx.py"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

CODE_BG = "F2F4F8"


def shade(par, color=CODE_BG):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def h1(t):
    doc.add_heading(t, level=1)


def h2(t):
    doc.add_heading(t, level=2)


def p(t, bold=False):
    par = doc.add_paragraph()
    run = par.add_run(t)
    run.bold = bold
    return par


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def numbered(items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def code(text, lang=""):
    if lang:
        par = doc.add_paragraph()
        r = par.add_run(lang)
        r.bold = True
        r.font.size = Pt(9)
    for line in text.strip("\n").split("\n"):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        run = par.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        shade(par)


def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        t.cell(0, i).text = htxt
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    doc.add_paragraph()
    return t


# ================= COVER =================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Cara Menjalankan Aplikasi\nFleetVision AI")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Lokal (Laptop) - HP Satu WiFi - Production\nhttps://dtrack.yusufandrika.dev")
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(0x47, 0x5B, 0x69)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Real-Time Fleet Tracking & Driver Monitoring System  |  Prototype / MVP")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

doc.add_page_break()

# ================= 1. GAMBARAN UMUM =================
h1("1. Gambaran Umum")
p("FleetVision AI terdiri dari 4 layanan yang berjalan bersamaan. "
  "Untuk menjalankan aplikasi, semuanya harus hidup (kecuali AI Service yang opsional bila hanya butuh GPS).")
table(["Layanan", "Perintah", "Port", "Fungsi"],
      [["Backend Laravel", "php artisan serve --host=0.0.0.0 --port=8000", "8000", "API + Dashboard web"],
       ["Reverb WebSocket", "php artisan reverb:start --host=0.0.0.0 --port=8080", "8080", "Update peta real-time"],
       ["AI Service (Python)", "python run.py  (folder ai-service)", "5000", "Deteksi seatbelt/fatigue/HP"],
       ["MySQL + Redis", "docker compose up -d mysql redis  (atau MySQL lokal)", "3306 / 6379", "Database + cache"]])

h2("1.1 Prasyarat")
bullets(["PHP 8.2+, Composer 2, Node.js 20+, Python 3.11+",
         "MySQL 8.0 (bisa via Docker) dan Redis (opsional di lokal)",
         "HP Android + Chrome (satu WiFi dengan laptop) untuk uji driver",
         "Koneksi internet untuk peta (Esri/OSM) dan rute OSRM"])

# ================= 2. LOKAL =================
h1("2. Menjalankan di Laptop (Localhost)")
h2("2.1 Setup pertama kali")
code("""
cd backend
composer install
cp .env.example .env
php artisan key:generate
php artisan migrate --seed
npm install
npm run build
""", lang="PowerShell (setup awal, sekali saja)")

h2("2.2 Menjalankan (4 terminal)")
code("""
# Terminal 1 - database (bila pakai Docker):
docker compose up -d mysql redis

# Terminal 2 - backend:
cd backend
php artisan serve --host=0.0.0.0 --port=8000

# Terminal 3 - realtime (opsional tapi disarankan):
cd backend
php artisan reverb:start --host=0.0.0.0 --port=8080

# Terminal 4 - AI service (opsional, wajib bila butuh deteksi kamera):
cd ai-service
python run.py
""", lang="PowerShell")
p("Buka http://localhost:8000/track untuk halaman publik, http://localhost:8000/login untuk masuk dashboard.",
  bold=True)

h2("2.3 Akun demo (hasil migrate --seed)")
table(["Role", "Email", "Password"],
      [["Admin", "admin@example.com", "password"],
       ["Driver", "budi@example.com", "password"],
       ["Driver", "maria.garcia@example.com", "password"],
       ["Driver", "ahmed.hassan@example.com", "password"],
       ["Driver", "sarah.johnson@example.com", "password"],
       ["Driver", "carlos.mendez@example.com", "password"]])

# ================= 3. HP SATU WIFI =================
h1("3. Menjalankan dari HP (Satu WiFi dengan Laptop)")
p("HP dipakai sebagai GPS Tracker + Kamera sopir. Tidak perlu install APK: cukup buka Chrome HP ke alamat laptop.")
numbered(["Pastikan HP dan laptop tersambung ke WiFi yang SAMA.",
          "Cari IP laptop (PowerShell):  ipconfig  -> catat IPv4 Wi-Fi, contoh 192.168.1.10.",
          "Edit backend/.env: APP_URL=http://192.168.1.10:8000, REVERB_HOST=192.168.1.10, "
          "VITE_REVERB_HOST=192.168.1.10, lalu jalankan npm run build.",
          "Jalankan serve dengan --host=0.0.0.0 (lihat Bab 2.2) dan izinkan port 8000/8080/5000 di Firewall Windows.",
          "Di Chrome HP buka http://192.168.1.10:8000/login, login sebagai driver, tekan Mulai Trip, "
          "izinkan Lokasi + Kamera."])
p("Catatan Chrome: GPS/kamera wajib konteks aman (HTTPS atau localhost). Bila muncul "
  "\"Only secure origins are allowed\", aktifkan flag chrome://flags/#unsafely-treat-insecure-origin-as-secure "
  "untuk alamat IP tersebut, atau sambungkan HP via USB + adb reverse agar bisa buka http://localhost:8000. "
  "Versi paling mudah tanpa flag adalah membuka https://dtrack.yusufandrika.dev (Bab 4).", bold=True)

# ================= 4. PRODUCTION =================
h1("4. Menjalankan di Production (dtrack.yusufandrika.dev)")
p("Aplikasi production sudah HTTPS sehingga GPS/kamera HP langsung jalan tanpa flag.")
numbered(["Masuk ke VPS, lalu:  git pull",
          "php artisan migrate --force  (wajib tiap ada perubahan database)",
          "npm run build  (wajib tiap ada perubahan tampilan; folder public/build tidak ikut git)",
          "php artisan optimize:clear  lalu  php artisan config:cache",
          "Restart ai-service bila ada perubahan deteksi AI.",
          "Buka https://dtrack.yusufandrika.dev/track (publik) dan /login (dashboard)."])
h2("4.1 Isi .env production (pokok)")
code("""
APP_URL=https://dtrack.yusufandrika.dev
SANCTUM_STATEFUL_DOMAINS=dtrack.yusufandrika.dev
SESSION_SECURE_COOKIE=true
VITE_REVERB_HOST=dtrack.yusufandrika.dev
VITE_REVERB_PORT=443
VITE_REVERB_SCHEME=https
VITE_AI_SERVICE_URL=https://dtrack.yusufandrika.dev/ai
VITE_AI_SERVICE_WS=wss://dtrack.yusufandrika.dev/ai/inference/stream
""", lang=".env (VPS)")

# ================= 5. ALUR PAKAI =================
h1("5. Alur Pakai per Peran")
h2("5.1 Admin")
numbered(["Login -> /admin (statistik), /admin/fleet (peta semua unit), /admin/fleet/{id} (detail + kamera live).",
          "Buat trip: /admin/trips -> Buat Trip -> pilih kendaraan + driver -> cari lokasi / klik peta "
          "untuk titik awal dan tujuan (estimasi jarak otomatis).",
          "Lihat history: /admin/trips -> tombol History per baris (atau panel otomatis di bawah tabel) -> "
          "peta rute rencana + jejak hijau aktual + Estimasi/Sudah jalan/Sisa/Durasi.",
          "Pantau alert: /admin/alerts."])
h2("5.2 Sopir")
numbered(["Login -> /driver. Bila ada Tugas rute dari admin, tekan Mulai Rute Ini; bila jalan bebas, tekan Mulai Trip.",
          "Selama jalan: GPS terkirim tiap 5 detik, kamera streaming ke AI tiap 500 ms. "
          "Kartu Jarak = jarak rencana admin; Ditempuh = jalan aktual; banner = Estimasi/Sisa/Ditempuh.",
          "Geser peta untuk melihat rute (mode Bebas); tekan Mengikuti untuk menempel ke kendaraan lagi.",
          "Selesaikan Trip untuk menutup. Riwayat ada di /driver/trips (klik untuk detail + peta)."])
h2("5.3 Publik (tanpa login)")
p("Buka /track, masukkan nomor plat (contoh B-1234-ABC): tampil posisi live, kecepatan, "
  "tujuan + estimasi + sisa bila kendaraan sedang bertrip.")

# ================= 6. VERIFIKASI =================
h1("6. Verifikasi Cepat")
code("""
# 1) API hidup + user ada:
curl -X POST https://dtrack.yusufandrika.dev/api/auth/login ^
 -H "Content-Type: application/json" -H "Accept: application/json" ^
 -d '{"email":"admin@example.com","password":"password"}'
# Harus: {"success":true,...}  (401 = user belum di-seed; 419 = CSRF, hubungi dev)

# 2) Cek user di server:
php artisan tinker --execute='echo App\\Models\\User::count();'

# 3) AI hidup:
curl https://dtrack.yusufandrika.dev/ai/inference/health
""", lang="Verifikasi")

# ================= 7. TROUBLESHOOTING =================
h1("7. Troubleshooting")
table(["Gejala", "Penyebab / Solusi"],
      [["419 CSRF token mismatch saat login", "Versi lama. Update kode terbaru (trustProxies, tanpa statefulApi), lalu optimize:clear + config:cache."],
       ["401 credentials incorrect", "User belum ada: jalankan php artisan db:seed --force. Akun benar: admin@example.com / password."],
       ["Only secure origins (HP)", "Buka via https://dtrack.yusufandrika.dev, atau aktifkan flag insecure-origin di Chrome HP."],
       ["Preview kamera admin hitam + Face tidak terdeteksi", "Sopir End Trip + Mulai lagi (stream lama macet). Pastikan preview di HP sopir terlihat."],
       ["History trip kosong (peta tanpa garis hijau)", "Trip PLANNED memang belum ada jejak. Trip lama kini diambil dari rentang waktu; bila tetap kosong berarti GPS tak pernah masuk (cek /admin/fleet)."],
       ["Peta bergeser sendiri saat di-scroll (driver)", "Itu mode Mengikuti. Geser peta untuk mode Bebas, tekan tombolnya untuk menempel lagi."],
       ["Garis rute lurus antar-pulau", "Wajar untuk penyeberangan laut (jalur feri). Di darat OSRM mengikuti jalan bila rute wajar."],
       ["Search PT tidak ketemu", "Data OSM gratis tidak selengkap Google. Coba tambah kota (cth: Dahana Subang) atau klik langsung di peta."],
       ["Halaman tidak berubah setelah update", "Wajib npm run build di mesin yang serving (VPS), lalu hard refresh Ctrl+Shift+R."]])

doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = end.add_run("— Selesai —  FleetVision AI")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

doc.save("Panduan-Menjalankan-Aplikasi-FleetVision-AI.docx")
print("saved Panduan-Menjalankan-Aplikasi-FleetVision-AI.docx")
