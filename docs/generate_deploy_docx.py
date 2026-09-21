"""Generate Panduan Deploy VPS XL FleetVision AI (.docx). Usage: python docs/generate_deploy_docx.py"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

CODE_BG = "F2F4F8"


def shade(par, color=CODE_BG):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def h1(t):
    doc.add_heading(t, level=1)


def h2(t):
    doc.add_heading(t, level=2)


def p(t, bold=False):
    par = doc.add_paragraph()
    run = par.add_run(t)
    run.bold = bold
    return par


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def numbered(items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def code(text, lang=""):
    if lang:
        par = doc.add_paragraph()
        r = par.add_run(lang)
        r.bold = True
        r.font.size = Pt(9)
    for line in text.strip("\n").split("\n"):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        run = par.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        shade(par)


def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        t.cell(0, i).text = htxt
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    doc.add_paragraph()
    return t


# ================= COVER =================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Panduan Deploy Produksi\nFleetVision AI — VPS XL")
r.bold = True
r.font.size = Pt(22)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Dari nol sampai online: Ubuntu + Docker + Nginx + SSL + Backup\nDisusun: September 2026")
r.font.size = Pt(11)

table(["Item", "Detail"], [
    ["Paket VPS", "XL — Rp 500.000/bln — 4 vCPU / 8 GB RAM / 160 GB SSD / Unlimited Traffic"],
    ["OS", "Ubuntu 22.04 LTS (tanpa cPanel/Plesk)"],
    ["Domain", "± Rp 150.000/tahun (.com)"],
    ["Estimasi total tahun pertama", "± Rp 6.150.000"],
])

# ================= 1 =================
h1("1. Rangkuman")
p("FleetVision AI terdiri dari 5 service yang wajib hidup 24 jam: backend Laravel (Nginx + PHP-FPM), "
  "MySQL 8, Redis 7, Laravel Reverb WebSocket, dan AI-service Python (YOLOv8 + MediaPipe). Perkiraan RAM: "
  "MySQL 1 GB + AI-service 2-3 GB + Laravel/Reverb/queue 1 GB + OS/Docker 1 GB = sekitar 6 GB. "
  "Karena itu paket XL 8 GB adalah pilihan yang pas (paket L 4 GB bisa jalan tapi mepet swap).")
p("Paket ini TIDAK BISA dipasang di shared hosting (termasuk paket Large Cloud Power), karena butuh akses root, "
  "Docker, proses persisten (reverb, queue worker, uvicorn), port custom 5000/8080, dan RAM 2-4 GB untuk model AI. "
  "Jangan ambil add-on cPanel/Plesk: berbayar ekstra dan memakan ±1 GB RAM secara sia-sia.", True)

# ================= 2 =================
h1("2. Prasyarat")
bullets([
    "VPS XL aktif + IP publik + password root (SSH).",
    "Domain (contoh: fleetvisionku.com) dengan akses DNS.",
    "Laptop dengan Git + SSH client (Windows: PowerShell/Terminal).",
    "Kode lokal sudah final dan ter-commit (lihat Bab 3).",
])

# ================= 3 =================
h1("3. Rapikan repo lokal dulu (wajib sebelum deploy)")
p("Repo sekarang kotor (sisa revert + file .pyc/.log). Bersihkan agar yang naik ke server steril:")
code("cd D:\\Code\\dtrack\n"
     "git status --short\n"
     "# pastikan file sampah tidak ikut: tambahkan ke .gitignore bila belum ada\n"
     "git add -A\n"
     'git commit -m "chore: siap deploy produksi"\n'
     "git push", "PowerShell (laptop)")
p("Isi .gitignore yang disarankan (append bila belum ada):")
code("__pycache__/\n*.pyc\n*.log\nbackend/.env\nai-service/.env\nbackend/node_modules/\npublic/build/", ".gitignore")

# ================= 4 =================
h1("4. Setup awal VPS (sekali saja)")
p("Login SSH lalu jalankan script berikut. Isinya: update OS, firewall (hanya buka 22/80/443), "
  "install Docker + Compose plugin, Node.js 20 (untuk build frontend), Nginx + Certbot, Git.")
code('ssh root@IP_VPS\n'
     "cat > /root/setup-server.sh <<'EOF'\n"
     "#!/bin/bash\nset -e\n"
     "apt update && apt upgrade -y\n"
     "apt install -y ca-certificates curl gnupg ufw nginx certbot python3-certbot-nginx git\n"
     "# --- firewall: JANGAN expose 5000/8080/3306/6379 ke publik ---\n"
     "ufw allow 22/tcp && ufw allow 80/tcp && ufw allow 443/tcp\n"
     "ufw --force enable\n"
     "# --- docker ---\n"
     "curl -fsSL https://get.docker.com | sh\n"
     "docker --version && docker compose version\n"
     "# --- node 20 (untuk npm run build di server) ---\n"
     "curl -fsSL https://deb.nodesource.com/setup_20.x | bash -\n"
     "apt install -y nodejs\n"
     "node -v\n"
     "mkdir -p /opt /opt/backups\n"
     'echo "SETUP SELESAI"\n'
     "EOF\nchmod +x /root/setup-server.sh\n/root/setup-server.sh", "SSH root")

# ================= 5 =================
h1("5. Arahkan domain (DNS)")
numbered([
    "Buka panel DNS domain, buat A record: Host @ → Value IP_VPS (TTL 300).",
    "Tunggu 5-30 menit, verifikasi dari laptop: nslookup fleetvisionku.com (ganti dengan domainmu).",
    "Semua contoh di dokumen ini memakai fleetvisionku.com — ganti dengan domain aslimu.",
])

# ================= 6 =================
h1("6. File konfigurasi produksi")
h2("6.1 docker-compose.prod.yml (BARU, di root repo)")
p("File ini melengkapi docker-compose.yml: menambah backend + reverb + queue-worker, "
  "mengunci port database ke localhost, dan mengganti password default. "
  "GANTI semua tulisan GANTI_PASSWORD_KUAT dan fleetvisionku.com:")
code('name: fleetvision-ai-prod\n\n'
     'services:\n'
     '  mysql:\n'
     '    ports:\n'
     '      - "127.0.0.1:3306:3306"\n'
     '    environment:\n'
     '      MYSQL_ROOT_PASSWORD: GANTI_PASSWORD_KUAT_1\n'
     '      MYSQL_DATABASE: fleetvision\n'
     '      MYSQL_USER: fleetvision\n'
     '      MYSQL_PASSWORD: GANTI_PASSWORD_KUAT_2\n\n'
     '  redis:\n'
     '    ports:\n'
     '      - "127.0.0.1:6379:6379"\n'
     '    command: ["redis-server", "--appendonly", "yes"]\n\n'
     '  ai-service:\n'
     '    ports:\n'
     '      - "127.0.0.1:5000:5000"\n'
     '    environment:\n'
     '      BACKEND_API_URL: https://fleetvisionku.com/api\n'
     '      REDIS_URL: redis://redis:6379/0\n'
     '      STREAM_SOURCE: none\n'
     '      INFERENCE_INTERVAL: "0.5"\n'
     '      CONFIDENCE_THRESHOLD: "0.5"\n'
     '      LOG_LEVEL: INFO\n\n'
     '  backend:\n'
     '    build:\n'
     '      context: ./backend\n'
     '      dockerfile: Dockerfile\n'
     '    container_name: fleetvision-backend\n'
     '    depends_on:\n'
     '      mysql: { condition: service_healthy }\n'
     '      redis: { condition: service_healthy }\n'
     '    env_file: [./backend/.env]\n'
     '    ports: ["127.0.0.1:8080:80"]\n'
     '    volumes: [backend_storage:/var/www/storage]\n'
     '    restart: unless-stopped\n\n'
     '  reverb:\n'
     '    build:\n'
     '      context: ./backend\n'
     '      dockerfile: Dockerfile\n'
     '    container_name: fleetvision-reverb\n'
     '    depends_on: [redis]\n'
     '    env_file: [./backend/.env]\n'
     '    command: ["php", "artisan", "reverb:start", "--host=0.0.0.0", "--port=8080"]\n'
     '    ports: ["127.0.0.1:8081:8080"]\n'
     '    restart: unless-stopped\n\n'
     '  queue-worker:\n'
     '    build:\n'
     '      context: ./backend\n'
     '      dockerfile: Dockerfile\n'
     '    container_name: fleetvision-queue\n'
     '    depends_on: [redis, mysql]\n'
     '    env_file: [./backend/.env]\n'
     '    command: ["php", "artisan", "queue:work", "--tries=3", "--sleep=3", "--timeout=120"]\n'
     '    volumes: [backend_storage:/var/www/storage]\n'
     '    restart: unless-stopped\n\n'
     'volumes:\n'
     '  backend_storage:', "docker-compose.prod.yml")

h2("6.2 backend/.env produksi (di server)")
p("Salin dari .env.example lalu sesuaikan. Yang PALING penting: APP_URL https, DB/Redis menunjuk nama service, "
  "Reverb menunjuk domain + 443/https, dan VITE_* (dibaca saat npm run build):")
code("APP_NAME=FleetVisionAI\nAPP_ENV=production\nAPP_DEBUG=false\n"
     "APP_URL=https://fleetvisionku.com\n\n"
     "DB_CONNECTION=mysql\nDB_HOST=mysql\nDB_PORT=3306\nDB_DATABASE=fleetvision\n"
     "DB_USERNAME=fleetvision\nDB_PASSWORD=GANTI_PASSWORD_KUAT_2\n\n"
     "BROADCAST_DRIVER=reverb\nCACHE_DRIVER=redis\nQUEUE_CONNECTION=redis\n"
     "SESSION_DRIVER=redis\n\nREDIS_HOST=redis\nREDIS_PORT=6379\nREDIS_PASSWORD=null\n\n"
     "REVERB_APP_ID=fleetvision-reverb\nREVERB_APP_KEY=fleetvision-key\n"
     "REVERB_APP_SECRET=GANTI_SECRET_ACAK_PANJANG\n"
     "REVERB_HOST=fleetvisionku.com\nREVERB_PORT=443\nREVERB_SCHEME=https\n"
     "REVERB_ALLOWED_ORIGINS=fleetvisionku.com\n\n"
     "SANCTUM_STATEFUL_DOMAINS=fleetvisionku.com\n\n"
     "VITE_AI_SERVICE_URL=https://fleetvisionku.com/ai\n"
     "VITE_AI_SERVICE_WS=wss://fleetvisionku.com/ai/inference/stream\n"
     "VITE_REVERB_HOST=fleetvisionku.com\nVITE_REVERB_PORT=443\n"
     "VITE_LIVE_POLL_MS=600\nVITE_AI_SEND_MS=500", "backend/.env (server)")

h2("6.3 Matikan auto-seed berulang")
p("PERHATIAN: backend/docker/start.sh menjalankan php artisan db:seed --force SETIAP container start — "
  "di produksi ini akan menggandakan data. Jalankan perintah ini sekali di server (sebelum build):", True)
code('sed -i "s/^php artisan db:seed --force/# DISABLED PROD: php artisan db:seed --force/" backend/docker/start.sh\n'
     "grep -n \"db:seed\" backend/docker/start.sh", "bash (di /opt/dtrack)")

h2("6.4 Reverse proxy + SSL di host VPS")
p("Nginx di HOST (bukan container) menjadi pintu depan: menangani SSL, meneruskan / ke backend, "
  "/ai/ ke ai-service, dan /app/ (websocket Reverb) ke service reverb:")
code("cat > /etc/nginx/sites-available/fleetvision <<'EOF'\n"
     "upstream laravel { server 127.0.0.1:8080; }\n"
     "upstream ai_service { server 127.0.0.1:5000; }\n"
     "upstream reverb { server 127.0.0.1:8081; }\n\n"
     "server {\n"
     "  listen 80;\n"
     "  server_name fleetvisionku.com;\n"
     "  client_max_body_size 10M;\n\n"
     "  location /ai/ {\n"
     "    proxy_pass http://ai_service/;\n"
     "    proxy_set_header Host $host;\n"
     "    proxy_set_header X-Forwarded-Proto $scheme;\n"
     "    proxy_read_timeout 86400;\n"
     "    proxy_send_timeout 86400;\n"
     "    proxy_http_version 1.1;\n"
     '    proxy_set_header Upgrade $http_upgrade;\n'
     '    proxy_set_header Connection "upgrade";\n'
     "  }\n\n"
     "  location /app/ {\n"
     "    proxy_pass http://reverb;\n"
     "    proxy_http_version 1.1;\n"
     '    proxy_set_header Upgrade $http_upgrade;\n'
     '    proxy_set_header Connection "upgrade";\n'
     "    proxy_set_header Host $host;\n"
     "    proxy_read_timeout 86400;\n"
     "    proxy_send_timeout 86400;\n"
     "  }\n\n"
     "  location / {\n"
     "    proxy_pass http://laravel;\n"
     "    proxy_set_header Host $host;\n"
     "    proxy_set_header X-Real-IP $remote_addr;\n"
     "    proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;\n"
     "    proxy_set_header X-Forwarded-Proto $scheme;\n"
     "  }\n"
     "}\nEOF\n"
     "ln -sf /etc/nginx/sites-available/fleetvision /etc/nginx/sites-enabled/fleetvision\n"
     "nginx -t && systemctl reload nginx", "/etc/nginx/sites-available/fleetvision")

# ================= 7 =================
h1("7. Deploy (script utama)")
code('cat > /opt/deploy.sh <<\'EOF\'\n'
     "#!/bin/bash\nset -e\ncd /opt/dtrack\n"
     'echo "== 1/5 git pull =="\ngit pull\n'
     'echo "== 2/5 build frontend =="\ncd backend\nnpm ci\nnpm run build\ncd ..\n'
     'echo "== 3/5 up containers =="\n'
     "docker compose -f docker-compose.yml -f docker-compose.prod.yml up -d --build\n"
     "sleep 20\n"
     'echo "== 4/5 migrate =="\n'
     "docker exec fleetvision-backend php artisan migrate --force\n"
     "docker exec fleetvision-backend php artisan storage:link || true\n"
     'echo "== 5/5 smoke test =="\n'
     "curl -sf http://127.0.0.1:5000/inference/health | head -c 200; echo\n"
     "curl -sf http://127.0.0.1:8080/ -o /dev/null && echo \"backend OK\"\n"
     'echo "DEPLOY SELESAI"\n'
     "EOF\nchmod +x /opt/deploy.sh\n/opt/deploy.sh", "SSH (setelah clone ke /opt/dtrack)")
p("Clone awal (hanya sekali, sebelum deploy pertama):")
code("cd /opt && git clone <URL_REPO_KAMU> dtrack && cd dtrack\n"
     "cp backend/.env.example backend/.env   # lalu isi sesuai Bab 6.2 (nano backend/.env)", "SSH")
p("Seed database HANYA sekali di awal (jangan diulang, data jadi ganda):", True)
code("docker exec fleetvision-backend php artisan db:seed --force", "SSH")

# ================= 8 =================
h1("8. Pasang SSL gratis")
code("certbot --nginx -d fleetvisionku.com\n"
     "# pilih opsi 2 (redirect HTTP ke HTTPS) saat ditanya\n"
     "certbot renew --dry-run   # pastikan auto-renew jalan", "SSH")
p("Setelah ini situs wajib diakses via https://fleetvisionku.com.")

# ================= 9 =================
h1("9. Verifikasi end-to-end (checklist)")
numbered([
    "Buka https://fleetvisionku.com/login — tampil kartu login tengah (bukan kepotong).",
    "Login admin@fleetvision.ai / password → masuk /admin.",
    "Sebagai driver di HP: mulai trip → GPS live bergerak di peta admin.",
    "Live Camera tampil ±2 FPS tanpa flicker; panel AI berubah Live dalam ±2 detik.",
    "Peta ber-tile Esri (BUKAN watermark API KEY REQUIRED). Jika masih CARTO: Ctrl+Shift+R.",
    "Cek service: docker ps (6 container UP) dan titik kesehatan di bawah ini.",
])
code("docker ps --format 'table {{.Names}}\\t{{.Status}}'\n"
     "curl -s http://127.0.0.1:5000/inference/health\n"
     "docker logs --tail 20 fleetvision-ai\n"
     "docker logs --tail 20 fleetvision-queue", "SSH cek kesehatan")

# ================= 10 =================
h1("10. APK Android untuk produksi")
numbered([
    "Di mobile/app/build.gradle.kts baris 17-18 ganti http://10.0.2.2:8000 dan :5000 menjadi "
    "https://fleetvisionku.com/api/ dan https://fleetvisionku.com/ai/.",
    "AndroidManifest sudah INTERNET + usesCleartextTraffic (HTTPS tetap jalan).",
    "Build release APK di Android Studio (Build > Generate Signed App Bundle/APK) lalu distribusikan ke driver.",
])

# ================= 11 =================
h1("11. Backup otomatis mingguan")
code('cat > /opt/backup-db.sh <<\'EOF\'\n'
     "#!/bin/bash\nset -e\n"
     "F=/opt/backups/fleetvision-$(date +%F).sql.gz\n"
     "docker exec fleetvision-db mysqldump -u root -pGANTI_PASSWORD_KUAT_1 fleetvision | gzip > \"$F\"\n"
     "find /opt/backups -name 'fleetvision-*.sql.gz' -mtime +30 -delete\n"
     'echo "backup: $F"\n'
     "EOF\nchmod +x /opt/backup-db.sh\n"
     "(crontab -l 2>/dev/null; echo \"0 2 * * 0 /opt/backup-db.sh >> /var/log/fleetvision-backup.log 2>&1\") | crontab -", "SSH")

# ================= 12 =================
h1("12. Perintah operasional harian")
table(["Kebutuhan", "Perintah"], [
    ["Lihat log backend", "docker logs -f fleetvision-backend"],
    ["Lihat log AI", "docker logs -f fleetvision-ai"],
    ["Restart satu service", "docker compose -f docker-compose.yml -f docker-compose.prod.yml restart backend"],
    ["Update kode", "cd /opt/dtrack && git pull && /opt/deploy.sh"],
    ["Cek RAM/CPU", "docker stats --no-stream"],
    ["Masuk mysql", "docker exec -it fleetvision-db mysql -u root -p"],
])

# ================= 13 =================
h1("13. Troubleshooting")
table(["Gejala", "Penyebab umum → solusi"], [
    ["502 Bad Gateway", "Container backend belum ready → tunggu 30 dtk, cek docker ps/logs."],
    ["Kamera/AI unreachable", "VITE_AI_SERVICE_URL salah → harus https://domain/ai, lalu npm run build ulang."],
    ["WS/Realtime mati", "Service reverb/queue-worker tidak jalan → docker ps, cek REVERB_* di .env."],
    ["Peta watermark API KEY", "Bundle lama di browser → Ctrl+Shift+R; pastikan build terbaru ter-deploy."],
    ["Seed ganda", "start.sh masih auto-seed → ulangi Bab 6.3."],
    ["Certbot gagal", "DNS belum propagasi → tunggu, cek nslookup, ulangi certbot."],
    ["OOM / container mati", "Cek docker stats; AI butuh 2-3 GB — di XL 8 GB seharusnya aman."],
])

doc.save(r"D:\Code\dtrack\docs\Panduan-Deploy-VPS-XL-FleetVision-AI.docx")
print("saved")
